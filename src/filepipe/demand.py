"""Public demand view from the OSB monthly insolvency counts by Forward Sortation Area (FSA).

Source: Office of the Superintendent of Bankruptcy, "Insolvency statistics by FSA, monthly".
These are real public counts. No debtor level data is involved.
"""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from filepipe.config import CONFIGS, RAW, RESULTS, WORK

OSB_DIR = RAW / "osb"
MONTHS = {m: i for i, m in enumerate(["january", "february", "march", "april", "may", "june", "july", "august",
                                         "september", "october", "november", "december"], 1)}
PROVINCE = {"NFLD": "NL", "NS": "NS", "PEI": "PE", "NB": "NB", "QUE": "QC", "ONT": "ON", "MAN": "MB", "SASK": "SK",
            "ALTA": "AB", "BC": "BC", "SAS": "SK", "MB": "MB", "PE": "PE", "NL": "NL", "NWT": "NT", "NNVT": "NU", "YUKON": "YT"}
COLS = {"consumer_bankruptcies": 1, "cb_liabilities": 3, "consumer_proposals": 7, "cp_liabilities": 9}
HOLDOUT = 6


def _month(cell: str) -> date:
    m = re.search(r"Month:\s*([A-Za-z]+)\s+(\d{4})", cell or "")
    if not m:
        raise ValueError(f"no month in {cell!r}")
    return date(int(m.group(2)), MONTHS[m.group(1).lower()], 1)


def _num(v) -> float:
    """Cells are numbers, blanks, or a lone '.' or '-' for nothing."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace(",", "").replace("$", "").strip()
    try:
        return float(t)
    except ValueError:
        return 0.0


FSA = re.compile(r"[ABCEGHJ-NPRSTVXY]\d[A-Z]")
DIGIT_FIX = {"O": "0", "I": "1", "L": "1", "S": "5", "Z": "2", "B": "8"}
UNPARSED = "???"


def clean_fsa(code: str) -> str:
    """Codes are as typed on the Statement of Affairs, so a few read 'KOK' or 'K0a'. Upper case, fix a letter
    where a digit belongs, and keep anything still invalid as one unparsed bucket so totals still add up."""
    c = re.sub(r"\s", "", str(code)).upper()[:3]
    if len(c) == 3:
        c = c[0] + DIGIT_FIX.get(c[1], c[1]) + c[2]
    return c if FSA.fullmatch(c) else UNPARSED


def _province(sheet: str) -> str:
    key = re.split(r"[_ ]", sheet.strip())[0].upper()
    return PROVINCE.get(key, key)


def read_workbook(path: Path) -> tuple[pd.DataFrame, list[dict]]:
    """One row per FSA and month. The sheet's own Total row is kept aside to check the sum."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, checks = [], []
    for name in wb.sheetnames:
        if name.lower().startswith("notes"):
            continue
        ws = wb[name]
        prov, month, total = _province(name), None, None
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 2:
                month = _month(str(r[0]))
            if i < 10 or r[0] is None:
                continue
            code = str(r[0]).strip()
            vals = {k: _num(r[c]) for k, c in COLS.items()}
            if code.lower() == "total":
                total = vals
                break
            rows.append({"month": month, "province": prov, "fsa": clean_fsa(code), **vals})
        if total is not None:
            got = {k: sum(x[k] for x in rows if x["month"] == month and x["province"] == prov) for k in ("consumer_bankruptcies", "consumer_proposals")}
            checks.append({"sheet": name, "file": path.name, "month": month.isoformat(), "province": prov,
                           "ok": all(abs(got[k] - total[k]) < 0.5 for k in got)})
    return pd.DataFrame(rows), checks


def load(osb_dir: Path = OSB_DIR) -> tuple[pd.DataFrame, list[dict]]:
    frames, checks = [], []
    for p in sorted(osb_dir.glob("*.xlsx")):
        df, ch = read_workbook(p)
        frames.append(df)
        checks += ch
    df = pd.concat(frames, ignore_index=True)
    df = df.groupby(["month", "province", "fsa"], as_index=False).sum()
    df["consumer_insolvencies"] = df["consumer_bankruptcies"] + df["consumer_proposals"]
    return df.sort_values(["month", "province", "fsa"]).reset_index(drop=True), checks


def offices(path: Path = CONFIGS / "farber_office_fsas.csv") -> pd.DataFrame:
    return pd.read_csv(path)


def monthly(df: pd.DataFrame, by: list[str] | None = None, col: str = "consumer_insolvencies") -> pd.DataFrame:
    """Wide table: one column per series, one row per month, zeros for months with no filing."""
    by = by or []
    if by:
        w = df.pivot_table(index="month", columns=by, values=col, aggfunc="sum", fill_value=0)
    else:
        w = df.groupby("month")[[col]].sum()
    idx = pd.date_range(min(w.index), max(w.index), freq="MS").date
    return w.reindex(idx, fill_value=0)


def seasonal_naive(hist: pd.DataFrame, h: int) -> np.ndarray:
    """Same month last year."""
    return hist.iloc[-12:].iloc[:h].values if h <= 12 else np.tile(hist.iloc[-12:].values, (2, 1))[:h]


def seasonal_trend(hist: pd.DataFrame, h: int) -> np.ndarray:
    """Same month last year, scaled by how the last 12 months compare with the 12 before them.
    The ratio is shrunk toward 1 for small series so one busy month does not swing a tiny FSA."""
    last, prev = hist.iloc[-12:].sum().values, hist.iloc[-24:-12].sum().values
    ratio = np.where(prev > 0, last / np.maximum(prev, 1e-9), 1.0)
    weight = prev / (prev + 60.0)
    ratio = 1 + weight * (np.clip(ratio, 0.5, 2.0) - 1)
    return seasonal_naive(hist, h) * ratio


FORECASTERS = {"seasonal-naive": seasonal_naive, "seasonal-trend": seasonal_trend}


def score(actual: np.ndarray, pred: np.ndarray) -> dict:
    a, p = np.asarray(actual, float), np.asarray(pred, float)
    nz = a > 0
    return {"mape": round(float(np.mean(np.abs(a[nz] - p[nz]) / a[nz])), 4) if nz.any() else None,
            "wape": round(float(np.abs(a - p).sum() / max(a.sum(), 1e-9)), 4)}


def backtest(w: pd.DataFrame, h: int = HOLDOUT) -> dict:
    hist, test = w.iloc[:-h], w.iloc[-h:]
    return {name: score(test.values, f(hist, h)) for name, f in FORECASTERS.items()}


def forecast(w: pd.DataFrame, name: str, h: int = HOLDOUT) -> pd.DataFrame:
    last = pd.Timestamp(w.index[-1])
    idx = [(last + pd.DateOffset(months=i)).date() for i in range(1, h + 1)]
    return pd.DataFrame(FORECASTERS[name](w, h), index=idx, columns=w.columns)


def run() -> dict:
    df, checks = load()
    WORK.mkdir(parents=True, exist_ok=True)
    df.to_csv(WORK / "osb_fsa_monthly.csv", index=False)
    off = offices()
    office_fsas = set(off["fsa"])
    df["office_fsa"] = df["fsa"].isin(office_fsas)
    unparsed = float(df[df["fsa"] == UNPARSED]["consumer_insolvencies"].sum())
    ex_qc = df[df["province"] != "QC"]
    first, last = min(df["month"]), max(df["month"])

    nat = monthly(ex_qc)
    ont = monthly(ex_qc[ex_qc["province"] == "ON"])
    by_fsa = monthly(ex_qc[ex_qc["fsa"] != UNPARSED], ["fsa"])
    big = by_fsa.loc[:, by_fsa.iloc[-12:].sum() >= 60]  # FSAs with at least 5 filings a month on average
    bt = {"canada_ex_quebec": backtest(nat), "ontario": backtest(ont),
          "fsa_level": backtest(by_fsa), "fsa_level_60_plus_a_year": backtest(big)}
    best = min(FORECASTERS, key=lambda n: bt["canada_ex_quebec"][n]["mape"] + bt["ontario"][n]["mape"])
    best_fsa = min(FORECASTERS, key=lambda n: bt["fsa_level"][n]["wape"])

    fc_nat = forecast(nat, best)
    fc_fsa = forecast(by_fsa, best_fsa)
    ytd = ex_qc[ex_qc["month"] >= date(last.year, 1, 1)]
    tot = float(ytd["consumer_insolvencies"].sum())
    in_off = float(ytd[ytd["office_fsa"]]["consumer_insolvencies"].sum())
    top = (ytd.groupby(["fsa", "province"])["consumer_insolvencies"].sum().reset_index()
           .sort_values("consumer_insolvencies", ascending=False))
    top["office"] = top["fsa"].isin(office_fsas)
    gaps = top[~top["office"]].head(15)
    nxt = fc_fsa.sum().rename("forecast_next_6").reset_index()
    nxt["office"] = nxt["fsa"].isin(office_fsas)
    fc_fsa.reset_index(names="month").melt(id_vars="month", var_name="fsa", value_name="forecast").to_csv(
        WORK / "osb_fsa_forecast.csv", index=False)

    out = {
        "source_files": sorted(p.name for p in OSB_DIR.glob("*.xlsx")),
        "months": [first.isoformat(), last.isoformat()], "n_months": int(nat.shape[0]),
        "fsa_rows": int(df.shape[0]), "sheet_total_checks": len(checks),
        "sheet_total_checks_passed": sum(c["ok"] for c in checks),
        "backtest_holdout_months": HOLDOUT, "backtest": bt, "chosen_forecaster": best, "chosen_forecaster_fsa": best_fsa,
        "rows_with_unparsed_fsa_consumer_insolvencies": int(unparsed),
        "forecast_canada_ex_quebec": {k.isoformat(): round(float(v), 0) for k, v in fc_nat.iloc[:, 0].items()},
        "ytd_months": f"{date(last.year, 1, 1).isoformat()} to {last.isoformat()}",
        "ytd_consumer_insolvencies_ex_quebec": int(tot),
        "ytd_consumer_proposals_ex_quebec": int(ytd["consumer_proposals"].sum()),
        "ytd_proposal_share": round(float(ytd["consumer_proposals"].sum()) / tot, 4),
        "office_fsas_listed": int(off["fsa"].nunique()),
        "office_fsas_found_in_data": int(len(office_fsas & set(df["fsa"]))),
        "ytd_in_office_fsas": int(in_off), "ytd_share_in_office_fsas": round(in_off / tot, 4),
        "forecast_next_6_in_office_fsas": round(float(nxt[nxt["office"]]["forecast_next_6"].sum()), 0),
        "top_fsas_without_office": gaps.to_dict(orient="records"),
    }
    RESULTS.mkdir(parents=True, exist_ok=True)
    (RESULTS / "demand.json").write_text(json.dumps(out, indent=1, default=str))
    return out


if __name__ == "__main__":
    r = run()
    print(json.dumps({k: v for k, v in r.items() if k not in ("top_fsas_without_office",)}, indent=1, default=str))
    print(r["top_fsas_without_office"][:5])
